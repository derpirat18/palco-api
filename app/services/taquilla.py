"""Taquilla (TQ): vender en el lugar, con caja, folios y corte a ciegas.

Reglas (decididas con JuanMa el 30 oct):

- **Quien**: PSD y el organizador hacen todo (son "jefe"). Su personal (cuentas
  de acceso) se asigna por evento como cajero o jefe de taquilla.
- **Caja**: la abre el jefe para un cajero, con fondo en efectivo y, si se
  entrega medio fisico (pulsera, gafete…), su rango de folios. Un cajero tiene
  una sola caja sin cerrar por evento. Solo el cajero vende en su caja.
- **Venta**: del mismo mapa que la venta en linea (lugares disponibles o con
  apartado vencido; nunca bloqueados), con el mismo candado: un UPDATE
  condicionado y el indice unico de partidas. Siempre con correo: los boletos
  llegan con su QR, igual que en linea. Formas de pago: efectivo (con cambio),
  tarjeta en terminal (con su autorizacion) y cortesia (con motivo).
- **Medio fisico**: si el cliente no tiene internet para abrir su correo y va
  a entrar ya, se le entrega pulsera o gafete. La venta guarda sus folios y
  sus boletos quedan marcados como "ya entro" en ese momento.
- **Correccion**: no hay cancelaciones por el cliente. Si el cajero se
  equivoco, el jefe anula la venta el mismo dia y antes del corte, con
  motivo; los lugares regresan a la venta.
- **Corte a ciegas**: el cajero cuenta efectivo, vouchers y folios que le
  sobran sin ver lo que dice el sistema; el jefe cierra viendo la diferencia
  y, si la hay, deja una nota.

En espera de Nadia (no se calcula aqui): quien se queda el efectivo, la
comision de PSD sobre la taquilla y el texto legal definitivo. La orden
congela la comision del organizador como cualquier venta.
"""

from __future__ import annotations

import io
import secrets
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import Integer, and_, cast, func, or_, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import (
    EVENTO_NO_EXISTE,
    FUNCION_NO_EXISTE,
    ErrorDeDominio,
    RecursoNoEncontrado,
)
from app.core.telefonos import normalizar_telefono
from app.models.base import (
    EstadoAsiento,
    EstadoBoleto,
    EstadoEvento,
    EstadoFuncion,
    EstadoOrden,
    RolUsuario,
)
from app.models.catalog import Client, Event, Performance, Seat, Venue, Zone
from app.models.inventory import InventoryMove, PerformanceSeat
from app.models.outbox import OutboxEvent
from app.models.sales import Cortesia, Customer, Order, OrderItem, Scan, Ticket
from app.models.taquilla import (
    ABIERTA,
    CERRADA,
    CONTADA,
    CORTESIA,
    EFECTIVO,
    JEFE,
    TERMINAL,
    CajaDeTaquilla,
    PersonalDeTaquilla,
    VentaDeTaquilla,
)
from app.models.users import User
from app.services.checkout import CERO, nuevo_codigo_de_boleto, nuevo_folio
from app.services.consentimientos import ConsentimientoService
from app.services.escaner import VALIDO, horario_de_escaneo
from app.services.tipos import ADULTO, asiento_con_tipo, asiento_de, tarifas_de
from app.services.venta import iva_de

CANAL = "taquilla"
CANAL_CORTESIA = "cortesia"
MAX_LUGARES = 20
CAJA_NO_EXISTE = "No encontramos esta caja."
VENTA_NO_EXISTE = "No encontramos esta venta."
_FUNCION_CERRADA = (EstadoFuncion.CANCELADA.value, EstadoFuncion.REALIZADA.value)


def reserva_de(caja: CajaDeTaquilla) -> str:
    """El apartado con que la caja guarda su reserva para vender sin red."""
    return f"caja-{caja.uuid.hex}"


class TaquillaInvalida(ErrorDeDominio):
    codigo = "taquilla_invalida"
    http_status = 409


class SinPermisoDeTaquilla(ErrorDeDominio):
    codigo = "permiso_insuficiente"
    http_status = 403


# --------------------------------------------------------------------------
# Quien
# --------------------------------------------------------------------------


async def puesto_en(db: AsyncSession, usuario: User, evento: Event) -> str | None:
    """jefe, cajero o None. PSD y el organizador del evento son jefe."""
    if usuario.rol == RolUsuario.ADMIN_PSD:
        return JEFE
    if usuario.client_id != evento.client_id:
        return None
    if usuario.rol == RolUsuario.ORGANIZADOR:
        return JEFE
    if usuario.rol == RolUsuario.ACCESO:
        return await db.scalar(
            select(PersonalDeTaquilla.puesto).where(
                PersonalDeTaquilla.user_id == usuario.id,
                PersonalDeTaquilla.event_id == evento.id,
            )
        )
    return None


async def evento_de_taquilla(
    db: AsyncSession, event_id: int, usuario: User, jefe: bool = False
) -> tuple[Event, str]:
    """El evento y el puesto de quien pregunta. Sin puesto, el evento no
    existe para el (igual que un recurso ajeno en el resto del panel)."""
    evento = await db.get(Event, event_id)
    puesto = await puesto_en(db, usuario, evento) if evento else None
    if evento is None or puesto is None:
        raise RecursoNoEncontrado(EVENTO_NO_EXISTE)
    if jefe and puesto != JEFE:
        raise SinPermisoDeTaquilla("Esto lo hace el jefe de taquilla.")
    return evento, puesto


def _ahora_local(evento_tz: str) -> datetime:
    return datetime.now(ZoneInfo(evento_tz))


def _dinero(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"))


# --------------------------------------------------------------------------
# Datos de la caja
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Totales:
    ventas: int
    boletos: int
    efectivo: Decimal
    terminal: Decimal
    cortesias: int
    folios_entregados: int
    anuladas: int
    efectivo_esperado: Decimal
    folios_esperados: int | None


def _folios_del_rango(caja: CajaDeTaquilla) -> int | None:
    if caja.folio_desde is None:
        return None
    return caja.folio_hasta - caja.folio_desde + 1


class TaquillaService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    # ------------------------------------------------------------------
    # Personal y configuracion
    # ------------------------------------------------------------------

    async def personal(self, evento: Event) -> list[tuple[PersonalDeTaquilla, User]]:
        return list(
            (
                await self.db.execute(
                    select(PersonalDeTaquilla, User)
                    .join(User, User.id == PersonalDeTaquilla.user_id)
                    .where(PersonalDeTaquilla.event_id == evento.id)
                    .order_by(User.nombre)
                )
            ).all()
        )

    async def asignar(
        self, evento: Event, cuenta_id: int, puesto: str, actor: User
    ) -> None:
        cuenta = await self.db.get(User, cuenta_id)
        if (
            cuenta is None
            or cuenta.rol != RolUsuario.ACCESO
            or cuenta.client_id != evento.client_id
            or not cuenta.activo
        ):
            raise TaquillaInvalida(
                "Solo se asignan cuentas de personal (control de acceso) del "
                "organizador del evento.",
                {"campo": "cuenta_id"},
            )
        await self.db.execute(
            pg_insert(PersonalDeTaquilla)
            .values(
                user_id=cuenta.id,
                event_id=evento.id,
                puesto=puesto,
                asignado_por=actor.id,
            )
            .on_conflict_do_update(
                constraint="uq_box_office_staff_user_event",
                set_={"puesto": puesto, "asignado_por": actor.id},
            )
        )
        await self.db.flush()

    async def quitar(self, evento: Event, cuenta_id: int) -> None:
        abierta = await self.db.scalar(
            select(CajaDeTaquilla.id).where(
                CajaDeTaquilla.event_id == evento.id,
                CajaDeTaquilla.cajero_id == cuenta_id,
                CajaDeTaquilla.estado != CERRADA,
            )
        )
        if abierta:
            raise TaquillaInvalida(
                "Tiene una caja sin cerrar: ciérrala antes de quitarlo de la taquilla."
            )
        fila = await self.db.scalar(
            select(PersonalDeTaquilla).where(
                PersonalDeTaquilla.event_id == evento.id,
                PersonalDeTaquilla.user_id == cuenta_id,
            )
        )
        if fila:
            await self.db.delete(fila)
            await self.db.flush()

    # ------------------------------------------------------------------
    # Cajas
    # ------------------------------------------------------------------

    async def abrir(
        self,
        evento: Event,
        cajero_id: int,
        fondo: Decimal,
        actor: User,
        medio: str | None = None,
        folio_desde: int | None = None,
        folio_hasta: int | None = None,
    ) -> CajaDeTaquilla:
        cajero = await self.db.get(User, cajero_id)
        if cajero is None or not cajero.activo:
            raise TaquillaInvalida("Esa cuenta no existe o está desactivada.")
        if await puesto_en(self.db, cajero, evento) is None:
            raise TaquillaInvalida(
                "Esa persona no está asignada a la taquilla de este evento.",
                {"campo": "cajero_id"},
            )
        if (folio_desde is None) != (folio_hasta is None):
            raise TaquillaInvalida(
                "Escribe el primer y el último folio del medio físico, o ninguno.",
                {"campo": "folio_hasta"},
            )
        if folio_desde is not None:
            if folio_desde < 1 or folio_hasta < folio_desde:
                raise TaquillaInvalida(
                    "El rango de folios no es válido: el último debe ser mayor o "
                    "igual que el primero.",
                    {"campo": "folio_hasta"},
                )
            if folio_hasta - folio_desde >= 5000:
                raise TaquillaInvalida(
                    "Un rango de folios puede tener hasta 5,000.",
                    {"campo": "folio_hasta"},
                )
            cruce = await self.db.scalar(
                select(CajaDeTaquilla.id).where(
                    CajaDeTaquilla.event_id == evento.id,
                    CajaDeTaquilla.folio_desde <= folio_hasta,
                    CajaDeTaquilla.folio_hasta >= folio_desde,
                )
            )
            if cruce:
                raise TaquillaInvalida(
                    "Esos folios ya se entregaron a otra caja de este evento.",
                    {"campo": "folio_desde"},
                )
        caja = CajaDeTaquilla(
            event_id=evento.id,
            cajero_id=cajero.id,
            abierta_por=actor.id,
            fondo=_dinero(fondo),
            medio=(medio or "").strip() or None,
            folio_desde=folio_desde,
            folio_hasta=folio_hasta,
            estado=ABIERTA,
        )
        try:
            async with self.db.begin_nested():
                self.db.add(caja)
                await self.db.flush()
        except IntegrityError as exc:
            raise TaquillaInvalida(
                f"{cajero.nombre} ya tiene una caja sin cerrar en este evento."
            ) from exc
        await self.db.refresh(caja)
        return caja

    async def cajas(
        self, evento: Event, usuario: User, puesto: str
    ) -> list[tuple[CajaDeTaquilla, User]]:
        consulta = (
            select(CajaDeTaquilla, User)
            .join(User, User.id == CajaDeTaquilla.cajero_id)
            .where(CajaDeTaquilla.event_id == evento.id)
            .order_by(CajaDeTaquilla.abierta_en.desc())
        )
        if puesto != JEFE:
            consulta = consulta.where(CajaDeTaquilla.cajero_id == usuario.id)
        return list((await self.db.execute(consulta)).all())

    async def caja(self, caja_uuid, usuario: User) -> tuple[CajaDeTaquilla, Event, str]:
        """La caja, su evento y el puesto. El cajero solo ve la suya."""
        caja = await self.db.scalar(
            select(CajaDeTaquilla).where(CajaDeTaquilla.uuid == caja_uuid)
        )
        if caja is None:
            raise RecursoNoEncontrado(CAJA_NO_EXISTE)
        evento = await self.db.get(Event, caja.event_id)
        puesto = await puesto_en(self.db, usuario, evento)
        if puesto is None or (puesto != JEFE and caja.cajero_id != usuario.id):
            raise RecursoNoEncontrado(CAJA_NO_EXISTE)
        return caja, evento, puesto

    async def ventas_de(self, caja: CajaDeTaquilla) -> list[tuple]:
        """(venta, orden, partidas) de la caja, la mas reciente primero."""
        filas = (
            await self.db.execute(
                select(VentaDeTaquilla, Order)
                .join(Order, Order.id == VentaDeTaquilla.order_id)
                .where(VentaDeTaquilla.caja_id == caja.id)
                .order_by(VentaDeTaquilla.vendida_en.desc(), VentaDeTaquilla.id.desc())
            )
        ).all()
        resultado = []
        for venta, orden in filas:
            partidas = list(
                (
                    await self.db.scalars(
                        select(OrderItem)
                        .where(OrderItem.order_id == orden.id)
                        .order_by(OrderItem.id)
                    )
                ).all()
            )
            resultado.append((venta, orden, partidas))
        return resultado

    async def totales(self, caja: CajaDeTaquilla) -> Totales:
        vigentes = and_(
            VentaDeTaquilla.caja_id == caja.id, VentaDeTaquilla.anulada_en.is_(None)
        )
        fila = (
            await self.db.execute(
                select(
                    func.count(VentaDeTaquilla.id),
                    func.coalesce(
                        func.sum(VentaDeTaquilla.total).filter(
                            VentaDeTaquilla.forma_pago == EFECTIVO
                        ),
                        0,
                    ),
                    func.coalesce(
                        func.sum(VentaDeTaquilla.total).filter(
                            VentaDeTaquilla.forma_pago == TERMINAL
                        ),
                        0,
                    ),
                    func.count(VentaDeTaquilla.id).filter(
                        VentaDeTaquilla.forma_pago == CORTESIA
                    ),
                    func.coalesce(
                        func.sum(func.jsonb_array_length(VentaDeTaquilla.folios)), 0
                    ),
                ).where(vigentes)
            )
        ).one()
        ventas, efectivo, terminal, cortesias, folios = fila
        boletos = await self.db.scalar(
            select(func.count(OrderItem.id))
            .join(VentaDeTaquilla, VentaDeTaquilla.order_id == OrderItem.order_id)
            .where(vigentes, OrderItem.anulado.is_(False))
        )
        anuladas = await self.db.scalar(
            select(func.count(VentaDeTaquilla.id)).where(
                VentaDeTaquilla.caja_id == caja.id,
                VentaDeTaquilla.anulada_en.is_not(None),
            )
        )
        rango = _folios_del_rango(caja)
        return Totales(
            ventas=ventas,
            boletos=boletos or 0,
            efectivo=_dinero(efectivo),
            terminal=_dinero(terminal),
            cortesias=cortesias,
            folios_entregados=int(folios),
            anuladas=anuladas or 0,
            efectivo_esperado=_dinero(caja.fondo + _dinero(efectivo)),
            folios_esperados=None if rango is None else rango - int(folios),
        )

    async def contar(
        self,
        caja: CajaDeTaquilla,
        usuario: User,
        puesto: str,
        efectivo: Decimal,
        terminal: Decimal,
        folios_sobrantes: int | None,
    ) -> CajaDeTaquilla:
        """El cajero cuenta lo que tiene, sin ver lo que espera el sistema."""
        if puesto != JEFE and caja.cajero_id != usuario.id:
            raise RecursoNoEncontrado(CAJA_NO_EXISTE)
        if caja.estado != ABIERTA:
            raise TaquillaInvalida("Esta caja ya se contó.")
        if caja.folio_desde is not None and folios_sobrantes is None:
            raise TaquillaInvalida(
                "Cuenta también los folios que te sobran.",
                {"campo": "folios_sobrantes"},
            )
        caja.efectivo_contado = _dinero(efectivo)
        caja.terminal_contado = _dinero(terminal)
        caja.folios_sobrantes = (
            folios_sobrantes if caja.folio_desde is not None else None
        )
        caja.contada_en = func.now()
        caja.estado = CONTADA
        # Ya no vende: lo que quedaba en su reserva regresa a la venta.
        await self.soltar_reserva(caja, usuario)
        await self.db.flush()
        await self.db.refresh(caja)
        return caja

    async def cerrar(
        self, caja: CajaDeTaquilla, usuario: User, nota: str | None
    ) -> CajaDeTaquilla:
        """El jefe cierra: queda lo que esperaba el sistema junto a lo contado."""
        if caja.estado != CONTADA:
            raise TaquillaInvalida(
                "Primero el cajero cuenta su caja; luego se cierra."
                if caja.estado == ABIERTA
                else "Esta caja ya está cerrada."
            )
        t = await self.totales(caja)
        diferencia = (
            caja.efectivo_contado != t.efectivo_esperado
            or caja.terminal_contado != t.terminal
            or (
                t.folios_esperados is not None
                and caja.folios_sobrantes != t.folios_esperados
            )
        )
        nota = (nota or "").strip()
        if diferencia and len(nota) < 5:
            raise TaquillaInvalida(
                "Hay diferencia entre lo contado y lo que dice el sistema: "
                "escribe qué pasó para cerrar.",
                {"campo": "nota"},
            )
        caja.efectivo_esperado = t.efectivo_esperado
        caja.terminal_esperado = t.terminal
        caja.folios_esperados = t.folios_esperados
        caja.nota_de_cierre = nota[:300] or None
        caja.cerrada_por = usuario.id
        caja.cerrada_en = func.now()
        caja.estado = CERRADA
        await self.db.flush()
        await self.db.refresh(caja)
        return caja

    # ------------------------------------------------------------------
    # Vender
    # ------------------------------------------------------------------

    async def vender(
        self,
        caja: CajaDeTaquilla,
        usuario: User,
        performance_id: int,
        seat_ids: list[int],
        email: str,
        nombre: str,
        forma_pago: str,
        clave_cliente: str,
        acepta_aviso: bool,
        aviso_version: str,
        apellidos: str | None = None,
        telefono: str | None = None,
        recibido: Decimal | None = None,
        autorizacion: str | None = None,
        motivo_cortesia: str | None = None,
        folios: list[int] | None = None,
        acepta_marketing: bool = False,
        tipos: dict[int, str] | None = None,
        sin_red: bool = False,
        vendida_en: datetime | None = None,
    ) -> VentaDeTaquilla:
        """Vende en la caja. `sin_red`: la venta se hizo sin internet con la
        reserva de la caja y el medio fisico, y se sube ahora; `vendida_en`
        es la hora del equipo en que se cobro."""
        previa = await self.db.scalar(
            select(VentaDeTaquilla).where(
                VentaDeTaquilla.clave_cliente == clave_cliente
            )
        )
        if previa is not None:
            if previa.caja_id != caja.id:
                raise TaquillaInvalida("Esa venta es de otra caja.")
            return previa  # el equipo reintento: no se cobra dos veces
        if caja.cajero_id != usuario.id:
            raise SinPermisoDeTaquilla("Solo el cajero de esta caja vende en ella.")
        if caja.estado != ABIERTA:
            raise TaquillaInvalida(
                "Esta caja ya se contó: abre otra para seguir vendiendo."
            )
        if not acepta_aviso:
            raise TaquillaInvalida(
                "Antes de vender, el cliente debe aceptar el aviso de privacidad.",
                {"campo": "acepta_aviso"},
            )
        telefono = normalizar_telefono(telefono)
        folios = list(folios or [])
        if not seat_ids:
            raise TaquillaInvalida("Elige al menos un lugar.")
        if len(seat_ids) > MAX_LUGARES or len(set(seat_ids)) != len(seat_ids):
            raise TaquillaInvalida(
                f"Una venta puede tener hasta {MAX_LUGARES} lugares, sin repetir."
            )
        if forma_pago == CORTESIA and len((motivo_cortesia or "").strip()) < 3:
            raise TaquillaInvalida(
                "Escribe el motivo de la cortesía.", {"campo": "motivo_cortesia"}
            )
        if forma_pago == TERMINAL and len((autorizacion or "").strip()) < 3:
            raise TaquillaInvalida(
                "Escribe el número de autorización del voucher de la terminal.",
                {"campo": "autorizacion"},
            )

        funcion = await self.db.get(Performance, performance_id)
        evento = await self.db.get(Event, caja.event_id)
        if funcion is None or funcion.event_id != evento.id:
            raise RecursoNoEncontrado(FUNCION_NO_EXISTE)
        if funcion.estado in _FUNCION_CERRADA or not funcion.inventario_generado:
            raise TaquillaInvalida("Esta función no está a la venta.")
        if evento.estado != EstadoEvento.PUBLICADO.value:
            raise TaquillaInvalida("El evento no está publicado.")
        ahora = await self.db.scalar(select(func.now()))
        hora = ahora
        if sin_red:
            if not folios:
                raise TaquillaInvalida(
                    "Sin internet, el cliente entra con el medio físico: escribe "
                    "sus folios.",
                    {"campo": "folios"},
                )
            # La hora del equipo, acotada: no antes de abrir la caja ni en el futuro.
            hora = min(max(vendida_en or ahora, caja.abierta_en), ahora)
        limite = funcion.inicia_en.timestamp() + evento.taquilla_minutos_despues * 60
        if hora.timestamp() > limite:
            raise TaquillaInvalida(
                "La taquilla ya no vende para esta función: pasó su horario.",
                {"motivo": "fuera_de_horario"},
            )

        if folios:
            await self._revisar_folios(caja, evento, folios, len(seat_ids))
            abre, cierra = horario_de_escaneo(funcion)
            # Sin red ya entro: la puerta lo dejo pasar con su medio fisico.
            if not sin_red and not (abre <= ahora <= cierra):
                raise TaquillaInvalida(
                    "Todavía no abre el acceso de esta función: no se puede entregar "
                    "el medio físico para entrar ya. Véndele con su correo.",
                    {"campo": "folios"},
                )

        filas = (
            await self.db.execute(
                select(PerformanceSeat, Seat, Zone)
                .join(Seat, Seat.id == PerformanceSeat.seat_id)
                .join(Zone, Zone.id == Seat.zone_id)
                .where(
                    PerformanceSeat.performance_id == funcion.id,
                    PerformanceSeat.seat_id.in_(seat_ids),
                )
                .order_by(Zone.orden, Seat.fila, Seat.id)
            )
        ).all()
        if {s.id for _, s, _ in filas} != set(seat_ids):
            raise TaquillaInvalida("Hay lugares que no pertenecen a esta función.")
        tipos = dict(tipos or {})
        if set(tipos) - set(seat_ids):
            raise TaquillaInvalida("Hay tipos de boleto para lugares que no se venden.")

        es_cortesia = forma_pago == CORTESIA
        iva = iva_de(evento)
        # Tipos de boleto (TB): el mismo precio que en linea para cada tipo.
        tarifas = await tarifas_de(self.db, funcion.id, iva)
        tipo_de = {ps.id: tipos.get(seat.id, ADULTO) for ps, seat, _ in filas}
        precios = [
            CERO
            if es_cortesia
            else tarifas.precio(
                tipo_de[ps.id],
                zona.id,
                ps.precio_final,
                ps.precio_personalizado,
                f"{zona.nombre} {seat.fila}-{seat.numero}",
            )
            for ps, seat, zona in filas
        ]
        total = _dinero(sum(precios, CERO))
        if forma_pago == EFECTIVO and (recibido is None or _dinero(recibido) < total):
            raise TaquillaInvalida(
                f"El efectivo recibido no alcanza: el total es ${total:,.2f}.",
                {"campo": "recibido", "total": str(total)},
            )
        nuevo_estado = (
            EstadoAsiento.CORTESIA.value if es_cortesia else EstadoAsiento.VENDIDO.value
        )
        previos = {ps.id: ps.estado for ps, _, _ in filas}
        async with self.db.begin_nested():
            tomadas = set(
                (
                    await self.db.scalars(
                        update(PerformanceSeat)
                        .where(
                            PerformanceSeat.id.in_(list(previos)),
                            or_(
                                PerformanceSeat.estado
                                == EstadoAsiento.DISPONIBLE.value,
                                and_(
                                    PerformanceSeat.estado
                                    == EstadoAsiento.APARTADO.value,
                                    PerformanceSeat.hold_expira_en <= func.now(),
                                ),
                                # La reserva de esta misma caja (fase 3).
                                and_(
                                    PerformanceSeat.estado
                                    == EstadoAsiento.APARTADO.value,
                                    PerformanceSeat.hold_id == reserva_de(caja),
                                ),
                            ),
                        )
                        .values(
                            estado=nuevo_estado,
                            hold_id=None,
                            hold_expira_en=None,
                            version=PerformanceSeat.version + 1,
                        )
                        .returning(PerformanceSeat.id)
                        .execution_options(synchronize_session="fetch")
                    )
                ).all()
            )
            if tomadas != set(previos):
                raise TaquillaInvalida(
                    "Alguien acaba de tomar alguno de esos lugares. Actualiza el mapa "
                    "y elige otros.",
                    {"motivo": "lugar_ocupado"},
                )

            comprador = await self._persona(email, nombre, apellidos, telefono)
            cliente = await self.db.get(Client, evento.client_id)
            iva_total = _dinero(sum((iva.iva_contenido(p) for p in precios), CERO))
            orden = Order(
                folio=nuevo_folio(),
                customer_id=comprador.id,
                performance_id=funcion.id,
                estado=EstadoOrden.PAGADA.value,
                subtotal=total - iva_total,
                iva=iva_total,
                descuento=CERO,
                total=total,
                comision_pct_snap=CERO if es_cortesia else cliente.comision_pct,
                iva_pct_snap=iva.pct,
                idempotency_key=f"taquilla-{clave_cliente}",
                pagada_en=hora,
                # Ya cobrado en el lugar: el barrido no debe intentar capturarla.
                capturada_en=hora,
                canal=CANAL_CORTESIA if es_cortesia else CANAL,
                nombre_snap=nombre.strip()[:120] or None,
                apellidos_snap=(apellidos or "").strip()[:120] or None,
                email_snap=email.strip().lower(),
                telefono_snap=telefono,
                tarjeta_snap=(
                    "Efectivo"
                    if forma_pago == EFECTIVO
                    else f"Tarjeta en terminal · aut. {autorizacion.strip()[:20]}"
                    if forma_pago == TERMINAL
                    else None
                ),
            )
            self.db.add(orden)
            await self.db.flush()
            partidas = [
                OrderItem(
                    order_id=orden.id,
                    performance_seat_id=ps.id,
                    zona_snap=zona.nombre,
                    asiento_snap=asiento_con_tipo(
                        asiento_de(seat.fila, seat.numero, zona.general),
                        tipo_de[ps.id],
                    ),
                    tipo=tipo_de[ps.id],
                    precio_unitario=precio,
                    iva=iva.iva_contenido(precio),
                    descuento=CERO,
                    importe=precio,
                )
                for (ps, seat, zona), precio in zip(filas, precios, strict=True)
            ]
            self.db.add_all(partidas)
            await self.db.flush()
            boletos = [
                Ticket(
                    order_item_id=p.id,
                    holder_customer_id=comprador.id,
                    codigo=nuevo_codigo_de_boleto(),
                    semilla=secrets.token_bytes(32),
                    # Con medio fisico entra en este momento.
                    estado=(
                        EstadoBoleto.USADO.value
                        if folios
                        else EstadoBoleto.VALIDO.value
                    ),
                )
                for p in partidas
            ]
            self.db.add_all(boletos)
            await self.db.flush()
            if es_cortesia:
                self.db.add(
                    Cortesia(
                        order_id=orden.id,
                        motivo=f"Taquilla: {motivo_cortesia.strip()}"[:120],
                        otorgada_por=usuario.id,
                    )
                )
            self.db.add_all(
                InventoryMove(
                    performance_seat_id=ps_id,
                    estado_previo=previo,
                    estado_nuevo=nuevo_estado,
                    motivo=CANAL,
                    actor_tipo="user",
                    actor_id=usuario.id,
                    order_id=orden.id,
                )
                for ps_id, previo in previos.items()
            )
            medio = caja.medio or "medio físico"
            for b, folio in zip(boletos, folios, strict=False):
                self.db.add(
                    Scan(
                        ticket_id=b.id,
                        performance_id=funcion.id,
                        user_id=usuario.id,
                        codigo_leido=b.codigo,
                        resultado=VALIDO,
                        escaneado_en=hora if sin_red else datetime.now(UTC),
                        offline=sin_red,
                        client_scan_id=f"taquilla-{secrets.token_hex(12)}",
                        nota=f"Entró por taquilla: {medio} {folio}"[:200],
                    )
                )
            venta = VentaDeTaquilla(
                order_id=orden.id,
                caja_id=caja.id,
                vendida_por=usuario.id,
                forma_pago=forma_pago,
                total=total,
                recibido=_dinero(recibido) if forma_pago == EFECTIVO else None,
                cambio=_dinero(recibido) - total if forma_pago == EFECTIVO else None,
                autorizacion=(autorizacion or "").strip()[:40] or None
                if forma_pago == TERMINAL
                else None,
                folios=folios,
                entro_con_medio=bool(folios),
                clave_cliente=clave_cliente,
                sin_red=sin_red,
                vendida_en=hora,
            )
            self.db.add(venta)
            consentimientos = ConsentimientoService(self.db)
            await consentimientos.registrar(
                comprador.id, "operacion", True, CANAL, aviso_version, orden.id
            )
            if acepta_marketing:
                await consentimientos.registrar(
                    comprador.id, "marketing", True, CANAL, aviso_version, orden.id
                )
            await self.db.flush()
            await self._anunciar(
                orden, evento, funcion, comprador, boletos, es_cortesia
            )
        await self.db.refresh(venta)
        return venta

    # ------------------------------------------------------------------
    # Reserva para vender sin internet (fase 3)
    # ------------------------------------------------------------------

    async def reservar(
        self,
        caja: CajaDeTaquilla,
        usuario: User,
        performance_id: int,
        cantidad: int,
        zone_id: int | None = None,
    ) -> int:
        """Aparta lugares para que la caja venda sin internet.

        Solo con red y con la caja abierta. Sin zona, toma los de la zona
        que menos se vende y, en ella, las ultimas filas. Vencen al cerrar
        la venta en linea (luego la taquilla vende con el mapa) y se
        sueltan al contar la caja. Devuelve cuantos tiene la caja ahora para
        esa funcion.
        """
        if caja.estado != ABIERTA:
            raise TaquillaInvalida("Esta caja ya se contó.")
        if caja.folio_desde is None:
            raise TaquillaInvalida(
                "Sin internet el cliente entra con el medio físico: esta caja "
                "necesita rango de folios para tener reserva.",
                {"campo": "folio_desde"},
            )
        if not 1 <= cantidad <= 50:
            raise TaquillaInvalida("La reserva es de 1 a 50 lugares.")
        funcion = await self.db.get(Performance, performance_id)
        evento = await self.db.get(Event, caja.event_id)
        if funcion is None or funcion.event_id != evento.id:
            raise RecursoNoEncontrado(FUNCION_NO_EXISTE)
        if funcion.estado in _FUNCION_CERRADA or not funcion.inventario_generado:
            raise TaquillaInvalida("Esta función no está a la venta.")
        ahora = await self.db.scalar(select(func.now()))
        vence = funcion.venta_cierra_en or funcion.inicia_en
        if vence <= ahora:
            raise TaquillaInvalida(
                "La venta en línea de esta función ya cerró: la taquilla vende "
                "con el mapa y no necesita reserva.",
                {"motivo": "venta_cerrada"},
            )
        hold = reserva_de(caja)
        ya = await self._reservadas(caja, funcion.id)
        faltan = cantidad - len(ya)
        if faltan <= 0:
            return len(ya)

        libre = or_(
            PerformanceSeat.estado == EstadoAsiento.DISPONIBLE.value,
            and_(
                PerformanceSeat.estado == EstadoAsiento.APARTADO.value,
                PerformanceSeat.hold_expira_en <= func.now(),
            ),
        )
        if zone_id is None:
            # La zona que menos se vende: la de mayor proporcion libre.
            proporcion = (
                await self.db.execute(
                    select(
                        Seat.zone_id,
                        func.avg(cast(libre, Integer)).label("libre"),
                    )
                    .join(PerformanceSeat, PerformanceSeat.seat_id == Seat.id)
                    .where(
                        PerformanceSeat.performance_id == funcion.id,
                        PerformanceSeat.estado != EstadoAsiento.KILL.value,
                    )
                    .group_by(Seat.zone_id)
                    .order_by(func.avg(cast(libre, Integer)).desc())
                )
            ).all()
            zonas = [z for z, _ in proporcion]
        else:
            zonas = [zone_id]
        candidatas: list[int] = []
        for z in zonas:
            candidatas += (
                await self.db.scalars(
                    select(PerformanceSeat.id)
                    .join(Seat, Seat.id == PerformanceSeat.seat_id)
                    .where(
                        PerformanceSeat.performance_id == funcion.id,
                        Seat.zone_id == z,
                        libre,
                    )
                    .order_by(Seat.fila.desc(), Seat.id)
                    .limit(faltan - len(candidatas))
                )
            ).all()
            if len(candidatas) >= faltan:
                break
        if not candidatas:
            raise TaquillaInvalida(
                "No quedan lugares libres para apartar en esa zona.",
                {"campo": "zone_id"},
            )
        tomadas = (
            await self.db.scalars(
                update(PerformanceSeat)
                .where(PerformanceSeat.id.in_(candidatas), libre)
                .values(
                    estado=EstadoAsiento.APARTADO.value,
                    hold_id=hold,
                    hold_expira_en=vence,
                    version=PerformanceSeat.version + 1,
                )
                .returning(PerformanceSeat.id)
                .execution_options(synchronize_session=False)
            )
        ).all()
        self.db.add_all(
            InventoryMove(
                performance_seat_id=ps_id,
                estado_previo=EstadoAsiento.DISPONIBLE.value,
                estado_nuevo=EstadoAsiento.APARTADO.value,
                motivo="reserva de taquilla",
                actor_tipo="user",
                actor_id=usuario.id,
                hold_id=hold,
            )
            for ps_id in tomadas
        )
        await self.db.flush()
        return len(ya) + len(tomadas)

    async def _reservadas(self, caja: CajaDeTaquilla, performance_id: int | None):
        q = select(PerformanceSeat.id).where(
            PerformanceSeat.estado == EstadoAsiento.APARTADO.value,
            PerformanceSeat.hold_id == reserva_de(caja),
            PerformanceSeat.hold_expira_en > func.now(),
        )
        if performance_id is not None:
            q = q.where(PerformanceSeat.performance_id == performance_id)
        return list((await self.db.scalars(q)).all())

    async def reserva(self, caja: CajaDeTaquilla) -> list[dict]:
        """Lo que la caja tiene apartado, con lo necesario para venderlo sin
        red: lugar, precio y tipos de boleto. El equipo lo guarda."""
        evento = await self.db.get(Event, caja.event_id)
        iva = iva_de(evento)
        filas = (
            await self.db.execute(
                select(PerformanceSeat, Seat, Zone, Performance)
                .join(Seat, Seat.id == PerformanceSeat.seat_id)
                .join(Zone, Zone.id == Seat.zone_id)
                .join(Performance, Performance.id == PerformanceSeat.performance_id)
                .where(
                    PerformanceSeat.estado == EstadoAsiento.APARTADO.value,
                    PerformanceSeat.hold_id == reserva_de(caja),
                    PerformanceSeat.hold_expira_en > func.now(),
                )
                .order_by(Performance.inicia_en, Zone.orden, Seat.fila, Seat.id)
            )
        ).all()
        tarifas: dict[int, object] = {}
        r = []
        for ps, seat, zona, funcion in filas:
            if funcion.id not in tarifas:
                tarifas[funcion.id] = await tarifas_de(self.db, funcion.id, iva)
            t = {} if ps.precio_personalizado else tarifas[funcion.id].de_zona(zona.id)
            r.append(
                {
                    "performance_id": funcion.id,
                    "inicia_en": funcion.inicia_en,
                    "seat_id": seat.id,
                    "zona_id": zona.id,
                    "zona": zona.nombre,
                    "fila": seat.fila,
                    "numero": seat.numero,
                    "precio": ps.precio_final,
                    "precio_nino": t.get("nino"),
                    "precio_mayor": t.get("mayor"),
                    "vence_en": ps.hold_expira_en,
                }
            )
        return r

    async def soltar_reserva(
        self, caja: CajaDeTaquilla, usuario: User, performance_id: int | None = None
    ) -> int:
        """Regresa a la venta lo que la caja tenia apartado."""
        q = update(PerformanceSeat).where(
            PerformanceSeat.estado == EstadoAsiento.APARTADO.value,
            PerformanceSeat.hold_id == reserva_de(caja),
        )
        if performance_id is not None:
            q = q.where(PerformanceSeat.performance_id == performance_id)
        soltadas = (
            await self.db.scalars(
                q.values(
                    estado=EstadoAsiento.DISPONIBLE.value,
                    hold_id=None,
                    hold_expira_en=None,
                    version=PerformanceSeat.version + 1,
                )
                .returning(PerformanceSeat.id)
                .execution_options(synchronize_session=False)
            )
        ).all()
        self.db.add_all(
            InventoryMove(
                performance_seat_id=ps_id,
                estado_previo=EstadoAsiento.APARTADO.value,
                estado_nuevo=EstadoAsiento.DISPONIBLE.value,
                motivo="reserva de taquilla soltada",
                actor_tipo="user",
                actor_id=usuario.id,
                hold_id=reserva_de(caja),
            )
            for ps_id in soltadas
        )
        await self.db.flush()
        return len(soltadas)

    async def _revisar_folios(
        self, caja: CajaDeTaquilla, evento: Event, folios: list[int], lugares: int
    ) -> None:
        if caja.folio_desde is None:
            raise TaquillaInvalida(
                "Esta caja no entrega medio físico: no tiene rango de folios.",
                {"campo": "folios"},
            )
        if len(folios) != lugares or len(set(folios)) != len(folios):
            raise TaquillaInvalida(
                "Escribe un folio distinto por cada lugar.", {"campo": "folios"}
            )
        fuera = [f for f in folios if not caja.folio_desde <= f <= caja.folio_hasta]
        if fuera:
            raise TaquillaInvalida(
                f"Esos folios no son de tu caja ({caja.folio_desde} a "
                f"{caja.folio_hasta}): {', '.join(map(str, fuera))}.",
                {"campo": "folios", "folios": fuera},
            )
        usados = (
            await self.db.scalars(
                select(VentaDeTaquilla.folios)
                .join(CajaDeTaquilla, CajaDeTaquilla.id == VentaDeTaquilla.caja_id)
                .where(
                    CajaDeTaquilla.event_id == evento.id,
                    VentaDeTaquilla.anulada_en.is_(None),
                )
            )
        ).all()
        ya = sorted(set(folios) & {f for lista in usados for f in lista})
        if ya:
            raise TaquillaInvalida(
                f"Esos folios ya se entregaron: {', '.join(map(str, ya))}.",
                {"campo": "folios", "folios": ya},
            )

    async def _persona(
        self, email: str, nombre: str, apellidos: str | None, telefono: str | None
    ) -> Customer:
        """Un correo, una persona: a una que ya existe solo se le llena lo que falta."""
        correo = email.strip().lower()
        await self.db.execute(
            pg_insert(Customer)
            .values(
                email=correo,
                nombre=nombre.strip() or None,
                apellidos=(apellidos or "").strip() or None,
                telefono=telefono,
            )
            .on_conflict_do_nothing(index_elements=[Customer.email])
        )
        persona = await self.db.scalar(
            select(Customer)
            .where(Customer.email == correo)
            .execution_options(populate_existing=True)
        )
        if persona.nombre is None:
            persona.nombre = nombre.strip() or None
        if persona.apellidos is None:
            persona.apellidos = (apellidos or "").strip() or None
        if persona.telefono is None:
            persona.telefono = telefono
        await self.db.flush()
        return persona

    async def _anunciar(
        self,
        orden: Order,
        evento: Event,
        funcion: Performance,
        comprador: Customer,
        boletos: list[Ticket],
        cortesia: bool,
    ) -> None:
        """Al bandeja de salida: el mismo correo con QR que una compra en linea
        (o el de cortesia)."""
        cliente = await self.db.get(Client, evento.client_id)
        if cortesia:
            payload = {
                "orden_uuid": str(orden.uuid),
                "folio": orden.folio,
                "funcion_uuid": str(funcion.uuid),
                "boletos": [str(b.uuid) for b in boletos],
            }
            tipo = "cortesia.emitida"
        else:
            payload = {
                "orden_uuid": str(orden.uuid),
                "folio": orden.folio,
                "estado": orden.estado,
                "cliente_uuid": str(cliente.uuid),
                "evento_uuid": str(evento.uuid),
                "funcion_uuid": str(funcion.uuid),
                "comprador_uuid": str(comprador.uuid),
                "subtotal": str(orden.subtotal),
                "iva": str(orden.iva),
                "total": str(orden.total),
                "moneda": orden.moneda,
                "boletos": [str(b.uuid) for b in boletos],
                "canal": CANAL,
            }
            tipo = "orden.pagada"
        self.db.add(OutboxEvent(tipo=tipo, agregado_uuid=orden.uuid, payload=payload))
        await self.db.flush()

    # ------------------------------------------------------------------
    # Despues de vender
    # ------------------------------------------------------------------

    async def venta(
        self, venta_uuid, usuario: User
    ) -> tuple[VentaDeTaquilla, CajaDeTaquilla, Event, str]:
        venta = await self.db.scalar(
            select(VentaDeTaquilla).where(VentaDeTaquilla.uuid == venta_uuid)
        )
        if venta is None:
            raise RecursoNoEncontrado(VENTA_NO_EXISTE)
        try:
            caja, evento, puesto = await self.caja(
                (await self.db.get(CajaDeTaquilla, venta.caja_id)).uuid, usuario
            )
        except RecursoNoEncontrado as exc:
            raise RecursoNoEncontrado(VENTA_NO_EXISTE) from exc
        return venta, caja, evento, puesto

    async def _mismo_dia(self, venta: VentaDeTaquilla, evento: Event) -> bool:
        tz = await self.db.scalar(
            select(Venue.zona_horaria).where(Venue.id == evento.venue_id)
        )
        zona = ZoneInfo(tz)
        return venta.vendida_en.astimezone(zona).date() == _ahora_local(tz).date()

    async def anular(
        self,
        venta: VentaDeTaquilla,
        caja: CajaDeTaquilla,
        evento: Event,
        usuario: User,
        motivo: str,
    ) -> VentaDeTaquilla:
        """Correccion por error del cajero: el mismo dia, antes del corte."""
        motivo = (motivo or "").strip()
        if len(motivo) < 5:
            raise TaquillaInvalida(
                "Escribe el motivo de la corrección.", {"campo": "motivo"}
            )
        if venta.anulada_en is not None:
            raise TaquillaInvalida("Esta venta ya se corrigió.")
        if caja.estado != ABIERTA:
            raise TaquillaInvalida(
                "Esta caja ya se contó: las correcciones son antes del corte."
            )
        if not await self._mismo_dia(venta, evento):
            raise TaquillaInvalida("Solo se corrigen ventas del mismo día.")
        orden = await self.db.get(Order, venta.order_id, with_for_update=True)
        partidas = list(
            (
                await self.db.scalars(
                    select(OrderItem).where(
                        OrderItem.order_id == orden.id, OrderItem.anulado.is_(False)
                    )
                )
            ).all()
        )
        ps_ids = [p.performance_seat_id for p in partidas]
        regresados = (
            await self.db.scalars(
                update(PerformanceSeat)
                .where(
                    PerformanceSeat.id.in_(ps_ids),
                    PerformanceSeat.estado.in_(
                        (EstadoAsiento.VENDIDO.value, EstadoAsiento.CORTESIA.value)
                    ),
                )
                .values(
                    estado=EstadoAsiento.DISPONIBLE.value,
                    version=PerformanceSeat.version + 1,
                )
                .returning(PerformanceSeat.id, PerformanceSeat.estado)
                .execution_options(synchronize_session="fetch")
            )
        ).all()
        self.db.add_all(
            InventoryMove(
                performance_seat_id=ps_id,
                estado_previo=(
                    EstadoAsiento.CORTESIA.value
                    if orden.canal == CANAL_CORTESIA
                    else EstadoAsiento.VENDIDO.value
                ),
                estado_nuevo=EstadoAsiento.DISPONIBLE.value,
                motivo="correccion_taquilla",
                actor_tipo="user",
                actor_id=usuario.id,
                order_id=orden.id,
            )
            for ps_id in regresados
        )
        await self.db.execute(
            update(Ticket)
            .where(Ticket.order_item_id.in_([p.id for p in partidas]))
            .values(estado=EstadoBoleto.CANCELADO.value)
            .execution_options(synchronize_session=False)
        )
        await self.db.execute(
            update(OrderItem)
            .where(OrderItem.id.in_([p.id for p in partidas]))
            .values(anulado=True)
            .execution_options(synchronize_session=False)
        )
        orden.estado = EstadoOrden.CANCELADA.value
        orden.cerrada_motivo = "error_de_cajero"
        venta.anulada_en = func.now()
        venta.anulada_por = usuario.id
        venta.motivo_anulacion = motivo[:200]
        await self.db.flush()
        await self.db.refresh(venta)
        return venta

    async def cambiar_correo(
        self, venta: VentaDeTaquilla, caja: CajaDeTaquilla, evento: Event, email: str
    ) -> Order:
        """El cliente dicto mal su correo: el jefe lo corrige (mismo dia, antes
        del corte) y los boletos se le mandan al bueno."""
        if venta.anulada_en is not None:
            raise TaquillaInvalida("Esta venta se corrigió: ya no tiene boletos.")
        if caja.estado != ABIERTA or not await self._mismo_dia(venta, evento):
            raise TaquillaInvalida(
                "El correo se corrige el mismo día y antes del corte de la caja."
            )
        orden = await self.db.get(Order, venta.order_id, with_for_update=True)
        anterior = orden.customer_id
        persona = await self._persona(
            email, orden.nombre_snap or "", orden.apellidos_snap, None
        )
        orden.customer_id = persona.id
        orden.email_snap = persona.email
        await self.db.execute(
            update(Ticket)
            .where(
                Ticket.order_item_id.in_(
                    select(OrderItem.id).where(OrderItem.order_id == orden.id)
                ),
                Ticket.holder_customer_id == anterior,
            )
            .values(holder_customer_id=persona.id)
            .execution_options(synchronize_session=False)
        )
        await self.db.flush()
        return orden

    # ------------------------------------------------------------------
    # Recibo y corte
    # ------------------------------------------------------------------

    async def recibo(
        self, venta: VentaDeTaquilla, caja: CajaDeTaquilla, evento: Event
    ) -> dict:
        orden = await self.db.get(Order, venta.order_id)
        funcion = await self.db.get(Performance, orden.performance_id)
        recinto = await self.db.get(Venue, evento.venue_id)
        cajero = (
            await self.db.get(User, venta.vendida_por) if venta.vendida_por else None
        )
        partidas = (
            await self.db.scalars(
                select(OrderItem)
                .where(OrderItem.order_id == orden.id)
                .order_by(OrderItem.id)
            )
        ).all()
        return {
            "folio": orden.folio,
            "uuid": venta.uuid,
            "evento": evento.nombre,
            "recinto": recinto.nombre,
            "ciudad": recinto.ciudad,
            "zona_horaria": recinto.zona_horaria,
            "inicia_en": funcion.inicia_en,
            "lugares": [
                {"zona": p.zona_snap, "asiento": p.asiento_snap, "precio": p.importe}
                for p in partidas
            ],
            "total": venta.total,
            "iva": orden.iva,
            "iva_pct": orden.iva_pct_snap,
            "forma_pago": venta.forma_pago,
            "recibido": venta.recibido,
            "cambio": venta.cambio,
            "autorizacion": venta.autorizacion,
            "medio": caja.medio,
            "folios": venta.folios,
            "correo": orden.email_snap,
            "cajero": cajero.nombre if cajero else None,
            "vendida_en": venta.vendida_en,
            "anulada": venta.anulada_en is not None,
            "texto_legal": evento.texto_legal_medio,
        }

    async def corte_excel(self, caja: CajaDeTaquilla, evento: Event) -> bytes:
        t = await self.totales(caja)
        cajero = await self.db.get(User, caja.cajero_id)
        tz = await self.db.scalar(
            select(Venue.zona_horaria).where(Venue.id == evento.venue_id)
        )
        zona = ZoneInfo(tz)

        def local(m):
            return m.astimezone(zona).replace(tzinfo=None) if m else None

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Corte"
        filas = [
            ("Evento", evento.nombre),
            ("Cajero", cajero.nombre),
            ("Abierta", local(caja.abierta_en)),
            ("Estado", caja.estado),
            ("Fondo inicial", caja.fondo),
            ("Ventas", t.ventas),
            ("Boletos", t.boletos),
            ("Cortesías", t.cortesias),
            ("Ventas corregidas por error", t.anuladas),
            ("Efectivo vendido", t.efectivo),
            ("Efectivo esperado (fondo + ventas)", t.efectivo_esperado),
            ("Efectivo contado", caja.efectivo_contado),
            ("Terminal vendido", t.terminal),
            ("Terminal contado (vouchers)", caja.terminal_contado),
            ("Medio físico", caja.medio),
            (
                "Folios",
                f"{caja.folio_desde} a {caja.folio_hasta}"
                if caja.folio_desde
                else None,
            ),
            ("Folios entregados", t.folios_entregados),
            ("Folios que deben sobrar", t.folios_esperados),
            ("Folios contados", caja.folios_sobrantes),
            ("Cerrada", local(caja.cerrada_en)),
            ("Nota de cierre", caja.nota_de_cierre),
        ]
        for nombre, valor in filas:
            hoja.append([nombre, valor])
            hoja.cell(hoja.max_row, 1).font = Font(bold=True)
            if isinstance(valor, Decimal):
                hoja.cell(hoja.max_row, 2).number_format = '"$"#,##0.00'
            elif isinstance(valor, datetime):
                hoja.cell(hoja.max_row, 2).number_format = "dd/mm/yyyy hh:mm"
        hoja.column_dimensions["A"].width = 36
        hoja.column_dimensions["B"].width = 30

        detalle = libro.create_sheet("Ventas")
        detalle.append(
            [
                "Hora", "Folio", "Lugares", "Forma de pago", "Total", "Recibido",
                "Cambio", "Autorización", "Folios del medio", "Correo", "Estado",
                "Motivo de corrección",
            ]
        )  # fmt: skip
        for celda in detalle[1]:
            celda.font = Font(bold=True)
        for venta, orden, partidas in reversed(await self.ventas_de(caja)):
            detalle.append(
                [
                    local(venta.vendida_en),
                    orden.folio,
                    ", ".join(f"{p.zona_snap} {p.asiento_snap}" for p in partidas),
                    venta.forma_pago,
                    venta.total,
                    venta.recibido,
                    venta.cambio,
                    venta.autorizacion,
                    ", ".join(map(str, venta.folios)),
                    orden.email_snap,
                    "corregida" if venta.anulada_en else "vigente",
                    venta.motivo_anulacion,
                ]
            )
            detalle.cell(detalle.max_row, 1).number_format = "dd/mm/yyyy hh:mm"
            for col in (5, 6, 7):
                detalle.cell(detalle.max_row, col).number_format = '"$"#,##0.00'
        for col, ancho in zip(
            "ABCDEFGHIJKL",
            (16, 16, 40, 14, 12, 12, 12, 14, 18, 28, 12, 30),
            strict=True,
        ):
            detalle.column_dimensions[col].width = ancho
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    # ------------------------------------------------------------------
    # Lo que ve cada quien
    # ------------------------------------------------------------------

    async def eventos_para(self, usuario: User) -> list[tuple[Event, str, Venue]]:
        """Los eventos con funciones por venir donde esta persona tiene puesto."""
        consulta = (
            select(Event, Venue)
            .join(Venue, Venue.id == Event.venue_id)
            .where(
                Event.estado == EstadoEvento.PUBLICADO.value,
                Event.id.in_(
                    select(Performance.event_id).where(
                        Performance.inicia_en
                        > func.now() - func.make_interval(0, 0, 0, 1),
                        Performance.estado.not_in(_FUNCION_CERRADA),
                    )
                ),
            )
            .order_by(Event.nombre)
        )
        if usuario.rol == RolUsuario.ORGANIZADOR:
            consulta = consulta.where(Event.client_id == usuario.client_id)
        elif usuario.rol == RolUsuario.ACCESO:
            consulta = consulta.where(
                Event.id.in_(
                    select(PersonalDeTaquilla.event_id).where(
                        PersonalDeTaquilla.user_id == usuario.id
                    )
                )
            )
        resultado = []
        for evento, recinto in (await self.db.execute(consulta)).all():
            puesto = await puesto_en(self.db, usuario, evento)
            if puesto:
                resultado.append((evento, puesto, recinto))
        return resultado

    async def funciones_de(self, evento: Event) -> list[Performance]:
        return list(
            (
                await self.db.scalars(
                    select(Performance)
                    .where(
                        Performance.event_id == evento.id,
                        Performance.estado.not_in(_FUNCION_CERRADA),
                        Performance.inventario_generado.is_(True),
                        Performance.inicia_en
                        > func.now()
                        - func.make_interval(
                            0, 0, 0, 0, 0, evento.taquilla_minutos_despues
                        ),
                    )
                    .order_by(Performance.inicia_en)
                )
            ).all()
        )
