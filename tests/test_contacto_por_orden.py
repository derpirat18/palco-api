"""Contacto por compra (OB1).

El registro del comprador es uno por correo para toda la plataforma. Lo
que importa: cada compra muestra el correo y el telefono que se
escribieron en ella; un organizador nunca ve un telefono que la persona
dio en otra compra (y menos con otro organizador); la cortesia lleva solo
lo que escribio quien la dio.
"""

import pytest
from sqlalchemy import select

from app.models.sales import Customer, Order
from app.services.confirmacion import armar_confirmacion
from tests.test_catalog import API
from tests.test_checkout import orden_iniciada
from tests.test_cortesias import de, esc, otorgar, stripe_falso  # noqa: F401

pytestmark = pytest.mark.asyncio


async def pagar(servicio, stripe, iniciada) -> Order:
    stripe.autorizar(iniciada.orden.stripe_payment_intent_id)
    return await servicio.finalizar(iniciada.orden.id)


async def ventas(cliente_http, usuarios, pid, quien="org_a") -> list[dict]:
    r = await cliente_http.get(
        f"{API}/performances/{pid}/ventas", headers=de(usuarios, quien)
    )
    assert r.status_code == 200, r.text
    return r.json()["ventas"]


async def test_cada_compra_muestra_su_telefono(
    cliente_http,
    db,
    clientes,
    usuarios,
    hold_service,
    stripe_falso,  # noqa: F811
):
    esc1, _, s1, i1 = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id, telefono="55 1111 1111"
    )
    await pagar(s1, stripe_falso, i1)
    esc2, _, s2, i2 = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id, telefono="55 2222 2222"
    )
    await pagar(s2, stripe_falso, i2)

    primera = await ventas(cliente_http, usuarios, esc1["performance_id"])
    segunda = await ventas(cliente_http, usuarios, esc2["performance_id"])

    assert {v["telefono"] for v in primera} == {"+525511111111"}
    assert {v["telefono"] for v in segunda} == {"+525522222222"}
    # El registro no se cambio: un correo ajeno no le cambia nada a su dueño.
    registro = await db.scalar(
        select(Customer).where(Customer.email == "ana.lopez@correo.mx")
    )
    assert registro.telefono == "+525511111111"


async def test_otro_organizador_no_ve_el_telefono_de_otra_compra(
    cliente_http,
    db,
    clientes,
    usuarios,
    hold_service,
    stripe_falso,  # noqa: F811
):
    # Ana compra con el organizador A y da su teléfono.
    _, _, s1, i1 = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id, telefono="55 1234 5678"
    )
    await pagar(s1, stripe_falso, i1)
    # El organizador B le da una cortesía al mismo correo, sin teléfono.
    from tests.test_apartados import funcion_a_la_venta

    esc_b = await funcion_a_la_venta(db, clientes[1].id)
    r = await otorgar(
        cliente_http,
        usuarios,
        esc_b["performance_id"],
        [esc_b["seat_ids"][0]],
        quien="org_b",
        email="Ana.Lopez@Correo.MX",
        nombre="Ana invitada",
    )
    assert r.status_code == 201, r.text
    assert r.json()["telefono"] is None

    vistas = await ventas(cliente_http, usuarios, esc_b["performance_id"], "org_b")
    assert [(v["email"], v["telefono"], v["titular"]) for v in vistas] == [
        ("ana.lopez@correo.mx", None, "Ana invitada")
    ]
    excel = await cliente_http.get(
        f"{API}/performances/{esc_b['performance_id']}/ventas.xlsx",
        headers=de(usuarios, "org_b"),
    )
    assert b"5512345678" not in excel.content


async def test_cortesia_con_telefono_lo_guarda_y_lo_valida(
    cliente_http,
    db,
    esc,  # noqa: F811
    usuarios,
):
    pid, seats = esc["performance_id"], esc["seat_ids"]
    malo = await otorgar(cliente_http, usuarios, pid, [seats[0]], telefono="123")
    assert malo.status_code == 422
    # Un teléfono malo no se lleva ningún lugar.
    bueno = await otorgar(
        cliente_http, usuarios, pid, [seats[0]], telefono="(55) 8765-4321"
    )
    assert bueno.status_code == 201, bueno.text
    assert bueno.json()["telefono"] == "+525587654321"
    vistas = await ventas(cliente_http, usuarios, pid)
    assert vistas[0]["telefono"] == "+525587654321"


async def test_el_correo_de_los_boletos_va_al_de_la_compra(
    db,
    clientes,
    hold_service,
    stripe_falso,  # noqa: F811
):
    _, _, servicio, iniciada = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id
    )
    orden = await pagar(servicio, stripe_falso, iniciada)
    assert orden.email_snap == "ana.lopez@correo.mx"
    # Cuando haya cuentas, una compra podrá llevar otro correo que el del perfil.
    orden.email_snap = "otro@correo.mx"
    await db.flush()

    mensaje = await armar_confirmacion(db, orden.uuid)

    assert mensaje is not None
    assert mensaje.para == "otro@correo.mx"


async def test_el_motivo_no_puede_empezar_como_formula(
    cliente_http,
    esc,  # noqa: F811
    usuarios,
):
    pid, seats = esc["performance_id"], esc["seat_ids"]
    for motivo in ("=HIPERVINCULO()", "+52", "-x", "@algo", "  =con espacios"):
        r = await otorgar(cliente_http, usuarios, pid, [seats[0]], motivo=motivo)
        assert r.status_code == 422, motivo
        campos = {
            c["campo"]: c["mensaje"] for c in r.json()["error"]["detalle"]["campos"]
        }
        assert campos["motivo"] == "Empieza con una letra o un número.", campos
    r = await otorgar(cliente_http, usuarios, pid, [seats[0]], motivo="2 x 1 prensa")
    assert r.status_code == 201, r.text


async def test_el_excel_trae_precio_sin_iva_que_cuadra(
    cliente_http,
    db,
    clientes,
    usuarios,
    hold_service,
    stripe_falso,  # noqa: F811
):
    import io
    from decimal import Decimal

    from openpyxl import load_workbook

    esc_, _, servicio, iniciada = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id
    )
    await pagar(servicio, stripe_falso, iniciada)
    r = await cliente_http.get(
        f"{API}/performances/{esc_['performance_id']}/ventas.xlsx",
        headers=de(usuarios, "org_a"),
    )
    hoja = load_workbook(io.BytesIO(r.content)).active
    titulos = [c.value for c in hoja[1]]
    precio, descuento, iva, total = (
        titulos.index(t) + 1 for t in ("Precio sin IVA", "Descuento", "IVA", "Total")
    )
    for fila in range(2, hoja.max_row + 1):
        p, d, i, t = (
            Decimal(str(hoja.cell(fila, c).value))
            for c in (precio, descuento, iva, total)
        )
        assert p - d + i == t
        assert p < t
