"""Listado de ventas de una funcion (RP1).

Lo que importa: lo ven PSD y el organizador dueno, nunca el personal de
acceso; lleva correo y telefono, por eso cada vista y cada descarga queda
en la bitacora; y el Excel no convierte en formula ni en numero lo que
escribio un comprador.
"""

import io
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.acceso import ConsultaDeDatos
from app.services.boletos import contenido_del_qr
from app.services.ventas import (
    ListadoDeVentas,
    RenglonDeVenta,
    ResumenDeVentas,
    a_excel,
)
from tests.test_catalog import API, igual_que_inexistente
from tests.test_checkout import orden_iniciada
from tests.test_cortesias import de, esc, otorgar, revocar, stripe_falso  # noqa: F401
from tests.test_escaner import asignar, boletos_de, leer

pytestmark = pytest.mark.asyncio


async def ventas(cliente_http, usuarios, pid, quien="org_a", xlsx=False):
    return await cliente_http.get(
        f"{API}/performances/{pid}/ventas{'.xlsx' if xlsx else ''}",
        headers=de(usuarios, quien),
    )


async def bitacora(db, pid) -> list[tuple[int, str]]:
    return [
        (c.user_id, c.accion)
        for c in (
            await db.scalars(
                select(ConsultaDeDatos)
                .where(ConsultaDeDatos.performance_id == pid)
                .order_by(ConsultaDeDatos.id)
            )
        ).all()
    ]


async def test_venta_web_y_cortesia_con_datos_del_comprador(
    cliente_http,
    db,
    clientes,
    usuarios,
    hold_service,
    stripe_falso,  # noqa: F811
):
    esc_, _, servicio, iniciada = await orden_iniciada(
        db, hold_service, stripe_falso, clientes[0].id
    )
    stripe_falso.autorizar(iniciada.orden.stripe_payment_intent_id)
    orden = await servicio.finalizar(iniciada.orden.id)
    pid = esc_["performance_id"]
    cortesia = (
        await otorgar(cliente_http, usuarios, pid, [esc_["seat_ids"][-1]])
    ).json()

    r = await ventas(cliente_http, usuarios, pid)

    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["resumen"] == {
        "boletos": 2,
        "cortesias": 1,
        "importe": "1600.00",
        "iva": cuerpo["resumen"]["iva"],
        "canceladas": 0,
    }
    web = [v for v in cuerpo["ventas"] if v["folio"] == orden.folio]
    assert len(web) == 2
    assert web[0] | {"comprada_en": None, "asiento": None, "iva": None} == {
        "folio": orden.folio,
        "comprada_en": None,
        "zona": "Luneta",
        "asiento": None,
        "titular": "Ana Lopez",
        "precio": "800.00",
        "descuento": "0.00",
        "iva": None,
        "total": "800.00",
        "canal": "web",
        "motivo_cortesia": None,
        "estado": "pagada",
        "entro_en": None,
        "email": "ana.lopez@correo.mx",
        "telefono": "+525512345678",
    }
    regalo = next(v for v in cuerpo["ventas"] if v["folio"] == cortesia["folio"])
    assert (regalo["canal"], regalo["motivo_cortesia"], regalo["total"]) == (
        "cortesia",
        "Prensa",
        "0.00",
    )


async def test_dice_quien_ya_entro_y_lo_revocado(cliente_http, db, esc, usuarios):  # noqa: F811
    pid = esc["performance_id"]
    entra = (await otorgar(cliente_http, usuarios, pid, [esc["seat_ids"][0]])).json()
    quitada = (await otorgar(cliente_http, usuarios, pid, [esc["seat_ids"][1]])).json()
    assert (await revocar(cliente_http, usuarios, quitada["uuid"])).status_code == 200
    await asignar(cliente_http, usuarios, [esc["event_id"]])
    boleto = (await boletos_de(db, entra["uuid"]))[0]
    await leer(cliente_http, usuarios, pid, contenido_del_qr(boleto.uuid))

    cuerpo = (await ventas(cliente_http, usuarios, pid)).json()

    por_folio = {v["folio"]: v for v in cuerpo["ventas"]}
    assert por_folio[entra["folio"]]["entro_en"] is not None
    assert por_folio[quitada["folio"]]["estado"] == "revocada"
    assert (cuerpo["resumen"]["cortesias"], cuerpo["resumen"]["canceladas"]) == (1, 1)


async def test_lo_ven_psd_y_el_organizador_nunca_la_puerta(
    cliente_http,
    esc,  # noqa: F811
    usuarios,
):
    pid = esc["performance_id"]
    await asignar(cliente_http, usuarios, [esc["event_id"]])
    assert (await ventas(cliente_http, usuarios, pid, "admin")).status_code == 200
    assert (await ventas(cliente_http, usuarios, pid, "org_a")).status_code == 200
    # Aun asignada al evento: lo suyo es la lista de la puerta, sin contacto.
    assert (await ventas(cliente_http, usuarios, pid, "acceso")).status_code == 403
    assert (
        await ventas(cliente_http, usuarios, pid, "acceso", xlsx=True)
    ).status_code == 403
    igual_que_inexistente(
        await ventas(cliente_http, usuarios, pid, "org_b"),
        await ventas(cliente_http, usuarios, 999999, "org_b"),
    )
    igual_que_inexistente(
        await ventas(cliente_http, usuarios, pid, "org_b", xlsx=True),
        await ventas(cliente_http, usuarios, 999999, "org_b", xlsx=True),
    )


async def test_cada_vista_y_descarga_queda_en_la_bitacora(
    cliente_http,
    db,
    esc,  # noqa: F811
    usuarios,
):
    pid = esc["performance_id"]
    await otorgar(cliente_http, usuarios, pid, [esc["seat_ids"][0]])

    await ventas(cliente_http, usuarios, pid, "org_a")
    r = await ventas(cliente_http, usuarios, pid, "admin", xlsx=True)
    await ventas(cliente_http, usuarios, pid, "org_b")  # ajeno: no cuenta

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in r.headers["content-disposition"]
    assert r.headers["cache-control"] == "no-store"
    assert await bitacora(db, pid) == [
        (usuarios["org_a"].id, "ver"),
        (usuarios["admin"].id, "descargar"),
    ]
    hoja = load_workbook(io.BytesIO(r.content)).active
    titulos = [c.value for c in hoja[1]]
    assert titulos[:5] == ["Folio", "Fecha de compra", "Zona", "Asiento", "A nombre de"]
    assert hoja.max_row == 2
    assert hoja.cell(2, 14).value == "prensa@medio.mx"


async def test_el_excel_no_ejecuta_lo_que_escribio_el_comprador():
    renglon = RenglonDeVenta(
        folio="PAL-12345678",
        comprada_en=datetime(2026, 10, 22, 2, 30, tzinfo=UTC),
        zona="Luneta",
        asiento="Fila A · 1",
        titular='=HYPERLINK("http://malo","clic")',
        precio=Decimal("800.00"),
        descuento=Decimal("0"),
        iva=Decimal("110.34"),
        total=Decimal("800.00"),
        canal="web",
        motivo_cortesia=None,
        estado="pagada",
        entro_en=None,
        email="@SUM(1+1)@x.mx",
        telefono="+525512345678",
    )
    listado = ListadoDeVentas(
        evento="Circo",
        recinto="Carpa",
        zona_horaria="America/Mexico_City",
        inicia_en=datetime(2026, 10, 25, tzinfo=UTC),
        resumen=ResumenDeVentas(1, 0, Decimal("800"), Decimal("110.34"), 0),
        ventas=(renglon,),
    )
    hoja = load_workbook(io.BytesIO(a_excel(listado))).active
    titular, telefono = hoja.cell(2, 5), hoja.cell(2, 15)
    assert (titular.value, titular.data_type) == (renglon.titular, "s")
    assert (telefono.value, telefono.data_type) == ("+525512345678", "s")
    # La hora de compra en la del recinto (CDMX, UTC-6): 21/10 20:30.
    assert hoja.cell(2, 2).value == datetime(2026, 10, 21, 20, 30)
    assert hoja.cell(2, 9).value == 800


async def test_sin_sesion_no_hay_listado(cliente_http, esc):  # noqa: F811
    r = await cliente_http.get(f"{API}/performances/{esc['performance_id']}/ventas")
    assert r.status_code == 401
